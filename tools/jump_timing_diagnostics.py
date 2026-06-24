from __future__ import annotations

import argparse
import csv
import math
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Literal

from openpyxl import load_workbook


Mode = Literal[
    "current",
    "compensated",
    "onset",
    "touch_associated",
    "lift_associated",
    "associated",
]
PairingMode = Literal["raw_triplet", "production_equivalent"]


@dataclass(frozen=True)
class TimingConfig:
    cols: int = 96
    touch_ratio_threshold: float = 0.12
    touch_ratio_max: float = 0.38
    lift_ratio_threshold: float = 0.05
    confirm_samples: int = 2
    min_valid_cluster_length: int = 10
    gap_threshold: int = 1
    spacing_cm: float = 1.04
    touch_onset_min_leds: int = 2
    lift_onset_max_leds: int = 3
    confirm_window_s: float = 0.10
    raw_track_max_missing_frames: int = 2
    raw_track_max_edge_gap_led: int = 4
    raw_track_max_centroid_shift_led: float = 6.0
    raw_track_max_touch_candidate_age_s: float | None = None


@dataclass(frozen=True)
class FrameSample:
    index: int
    timestamp: float
    bits: list[int]


@dataclass(frozen=True)
class RawCluster:
    start: int
    end: int
    length: int
    centroid_idx: float
    centroid_cm: float


@dataclass(frozen=True)
class ClusterStats:
    active_led_count: int
    raw_clusters: tuple[RawCluster, ...]
    primary_raw_cluster_index: int | None
    raw_primary_cluster_length: int
    valid_primary_cluster_length: int
    ratio: float
    centroid_cm: float | None


@dataclass
class RawTrack:
    track_id: int
    first_seen_frame: int
    first_seen_time: float
    last_seen_frame: int
    last_seen_time: float
    last_start: int
    last_end: int
    last_centroid_idx: float
    max_length: int
    seen_count: int = 1
    miss_count: int = 0
    is_active: bool = True

    def update(self, frame: FrameSample, cluster: RawCluster) -> None:
        self.last_seen_frame = frame.index
        self.last_seen_time = frame.timestamp
        self.last_start = cluster.start
        self.last_end = cluster.end
        self.last_centroid_idx = cluster.centroid_idx
        self.max_length = max(self.max_length, cluster.length)
        self.seen_count += 1
        self.miss_count = 0

    def miss(self, config: TimingConfig) -> None:
        self.miss_count += 1
        if self.miss_count > config.raw_track_max_missing_frames:
            self.is_active = False


@dataclass(frozen=True)
class EventRecord:
    kind: str
    time: float
    ratio: float
    centroid_cm: float | None
    frame_index: int
    confirm_time: float
    first_confirm_frame_time: float
    onset_time: float | None
    associated_track_id: int | None = None
    boundary_time: float | None = None
    fallback_reason: str = ""


@dataclass(frozen=True)
class FrameTrace:
    variant: str
    confirm_window_s: float
    frame_index: int
    timestamp: float
    active_led_count: int
    raw_primary_cluster_length: int
    valid_primary_cluster_length: int
    ratio: float
    state: str
    touch_condition: bool
    lift_condition: bool
    touch_streak: int
    lift_streak: int
    event_kind: str
    event_time_current: float | None
    event_time_first_confirm_frame: float | None
    onset_time: float | None
    confirm_time: float | None
    associated_track_id: int | None
    boundary_time: float | None
    fallback_reason: str


@dataclass(frozen=True)
class DetectorResult:
    events: list[EventRecord]
    traces: list[FrameTrace]


@dataclass(frozen=True)
class JumpRow:
    jump_index: int
    touch_start_s: float | None
    lift_s: float | None
    landing_s: float | None
    contact_time_s: float | None
    air_time_s: float | None


@dataclass(frozen=True)
class ManualRow:
    jump_index: int
    contact_time_s: float
    air_time_s: float


def _hex_to_bits(hex_string: str, cols: int = 96) -> list[int]:
    bits: list[int] = []
    for token in hex_string.strip().split():
        value = int(token, 16)
        bits.extend((value >> bit_index) & 1 for bit_index in range(8))
    if len(bits) < cols:
        bits.extend([0] * (cols - len(bits)))
    return bits[:cols]


def iter_excel_led_frames(path: str | Path) -> Iterable[FrameSample]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["LED Frames"] if "LED Frames" in workbook.sheetnames else workbook.active
        header = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        try:
            timestamp_index = header.index("timestamp")
            hex_index = header.index("hex_string")
        except ValueError as exc:
            raise ValueError("Excel sheet must contain timestamp and hex_string columns") from exc

        frame_index = 0
        for row in sheet.iter_rows(min_row=2, values_only=True):
            timestamp = row[timestamp_index]
            hex_string = row[hex_index]
            if timestamp is None or hex_string is None:
                continue
            frame_index += 1
            yield FrameSample(
                index=frame_index,
                timestamp=float(timestamp),
                bits=_hex_to_bits(str(hex_string)),
            )
    finally:
        workbook.close()


def _extract_raw_clusters(bits: list[int], config: TimingConfig) -> tuple[RawCluster, ...]:
    active_indices = [index for index, value in enumerate(bits[: config.cols]) if value]
    if not active_indices:
        return ()

    ranges: list[tuple[int, int]] = []
    start = previous = active_indices[0]
    for index in active_indices[1:]:
        if index - previous > config.gap_threshold:
            ranges.append((start, previous))
            start = index
        previous = index
    ranges.append((start, previous))

    return tuple(
        RawCluster(
            start=start,
            end=end,
            length=end - start + 1,
            centroid_idx=(start + end) / 2.0,
            centroid_cm=((start + end) / 2.0) * config.spacing_cm,
        )
        for start, end in ranges
    )


def _cluster_stats(bits: list[int], config: TimingConfig) -> ClusterStats:
    raw_clusters = _extract_raw_clusters(bits, config)
    active_count = sum(cluster.length for cluster in raw_clusters)
    if not raw_clusters:
        return ClusterStats(0, (), None, 0, 0, 0.0, None)

    primary_index, primary_cluster = max(
        enumerate(raw_clusters),
        key=lambda item: (item[1].length, -item[1].start),
    )
    raw_length = primary_cluster.length
    if raw_length < config.min_valid_cluster_length:
        return ClusterStats(
            active_count,
            raw_clusters,
            primary_index,
            raw_length,
            0,
            0.0,
            None,
        )

    valid_length = raw_length
    return ClusterStats(
        active_led_count=active_count,
        raw_clusters=raw_clusters,
        primary_raw_cluster_index=primary_index,
        raw_primary_cluster_length=raw_length,
        valid_primary_cluster_length=valid_length,
        ratio=valid_length / config.cols,
        centroid_cm=primary_cluster.centroid_cm,
    )


def _decay_streak(streak: int, condition_times: list[float]) -> int:
    next_streak = max(0, streak - 1)
    while len(condition_times) > next_streak:
        condition_times.pop(0)
    return next_streak


def _edge_gap(track: RawTrack, cluster: RawCluster) -> int:
    if cluster.end < track.last_start:
        return track.last_start - cluster.end
    if cluster.start > track.last_end:
        return cluster.start - track.last_end
    return 0


def _overlap_length(track: RawTrack, cluster: RawCluster) -> int:
    start = max(track.last_start, cluster.start)
    end = min(track.last_end, cluster.end)
    return max(0, end - start + 1)


def _raw_track_match_score(
    track: RawTrack,
    cluster: RawCluster,
    config: TimingConfig,
) -> tuple[int, int, float, int, int] | None:
    overlap = _overlap_length(track, cluster)
    edge_gap = _edge_gap(track, cluster)
    centroid_shift = abs(cluster.centroid_idx - track.last_centroid_idx)

    if overlap > 0:
        match_type = 0
    elif edge_gap <= config.raw_track_max_edge_gap_led:
        match_type = 1
    elif (
        cluster.length > 1
        and centroid_shift <= config.raw_track_max_centroid_shift_led
    ):
        match_type = 2
    else:
        return None

    return (
        match_type,
        edge_gap,
        centroid_shift,
        track.miss_count,
        -overlap,
    )


def _new_raw_track(track_id: int, frame: FrameSample, cluster: RawCluster) -> RawTrack:
    return RawTrack(
        track_id=track_id,
        first_seen_frame=frame.index,
        first_seen_time=frame.timestamp,
        last_seen_frame=frame.index,
        last_seen_time=frame.timestamp,
        last_start=cluster.start,
        last_end=cluster.end,
        last_centroid_idx=cluster.centroid_idx,
        max_length=cluster.length,
    )


def _update_raw_tracks(
    tracks: list[RawTrack],
    next_track_id: int,
    frame: FrameSample,
    clusters: tuple[RawCluster, ...],
    config: TimingConfig,
) -> tuple[int, dict[int, int]]:
    active_tracks = [track for track in tracks if track.is_active]
    cluster_to_track: dict[int, int] = {}
    assigned_tracks: set[int] = set()
    assigned_clusters: set[int] = set()
    candidates: list[tuple[tuple[int, int, float, int, int], int, int, RawTrack]] = []

    for track in active_tracks:
        for cluster_index, cluster in enumerate(clusters):
            score = _raw_track_match_score(track, cluster, config)
            if score is not None:
                candidates.append((score, track.track_id, cluster_index, track))

    candidates.sort(
        key=lambda item: (
            item[0],
            item[1],
            clusters[item[2]].start,
            item[2],
        )
    )

    for _score, track_id, cluster_index, track in candidates:
        if track_id in assigned_tracks or cluster_index in assigned_clusters:
            continue
        track.update(frame, clusters[cluster_index])
        assigned_tracks.add(track_id)
        assigned_clusters.add(cluster_index)
        cluster_to_track[cluster_index] = track_id

    for track in active_tracks:
        if track.track_id not in assigned_tracks:
            track.miss(config)

    for cluster_index, cluster in enumerate(clusters):
        if cluster_index in assigned_clusters:
            continue
        track = _new_raw_track(next_track_id, frame, cluster)
        tracks.append(track)
        cluster_to_track[cluster_index] = next_track_id
        next_track_id += 1

    return next_track_id, cluster_to_track


def _find_raw_track(tracks: list[RawTrack], track_id: int | None) -> RawTrack | None:
    if track_id is None:
        return None
    for track in tracks:
        if track.track_id == track_id:
            return track
    return None


def _event_time_for_mode(
    mode: Mode,
    kind: str,
    confirm_time: float,
    first_confirm_frame_time: float,
    onset_time: float | None,
    boundary_time: float | None,
    confirm_window_s: float,
) -> float:
    if mode == "current":
        return confirm_time
    if mode == "compensated":
        return first_confirm_frame_time
    if mode == "associated" and boundary_time is not None:
        return boundary_time
    if mode == "touch_associated" and kind == "touch" and boundary_time is not None:
        return boundary_time
    if mode == "lift_associated" and kind == "lift" and boundary_time is not None:
        return boundary_time
    if mode in {"associated", "touch_associated", "lift_associated"}:
        return first_confirm_frame_time
    if onset_time is not None and confirm_time - onset_time <= confirm_window_s:
        return onset_time
    return first_confirm_frame_time


def run_detector(
    frames: Iterable[FrameSample],
    mode: Mode,
    config: TimingConfig | None = None,
) -> DetectorResult:
    if config is None:
        config = TimingConfig()
    if mode not in {
        "current",
        "compensated",
        "onset",
        "touch_associated",
        "lift_associated",
        "associated",
    }:
        raise ValueError(f"Unknown mode: {mode}")

    state = "air"
    touch_streak = 0
    lift_streak = 0
    touch_condition_times: list[float] = []
    lift_condition_times: list[float] = []
    touch_onset_time: float | None = None
    lift_onset_time: float | None = None
    raw_tracks: list[RawTrack] = []
    next_track_id = 0
    active_contact_track_id: int | None = None
    events: list[EventRecord] = []
    traces: list[FrameTrace] = []

    for frame in frames:
        stats = _cluster_stats(frame.bits, config)
        next_track_id, cluster_to_track = _update_raw_tracks(
            raw_tracks,
            next_track_id,
            frame,
            stats.raw_clusters,
            config,
        )
        primary_track_id = (
            cluster_to_track.get(stats.primary_raw_cluster_index)
            if stats.primary_raw_cluster_index is not None
            else None
        )
        state_before = state
        touch_condition = (
            state == "air"
            and stats.ratio >= config.touch_ratio_threshold
            and stats.ratio < config.touch_ratio_max
        )
        lift_condition = (
            state == "ground"
            and (
                stats.ratio <= config.lift_ratio_threshold
                or stats.valid_primary_cluster_length == 0
            )
        )

        if mode == "onset":
            if state == "air":
                touch_onset = (
                    stats.active_led_count >= config.touch_onset_min_leds
                    or stats.raw_primary_cluster_length >= config.touch_onset_min_leds
                )
                if touch_onset and touch_onset_time is None:
                    touch_onset_time = frame.timestamp
                if (
                    touch_onset_time is not None
                    and not touch_condition
                    and frame.timestamp - touch_onset_time > config.confirm_window_s
                ):
                    touch_onset_time = None
                    touch_streak = 0
                    touch_condition_times.clear()
            elif state == "ground":
                lift_onset = (
                    stats.active_led_count <= config.lift_onset_max_leds
                    or stats.raw_primary_cluster_length <= config.lift_onset_max_leds
                )
                if lift_onset and lift_onset_time is None:
                    lift_onset_time = frame.timestamp
                if (
                    lift_onset_time is not None
                    and not lift_condition
                    and frame.timestamp - lift_onset_time > config.confirm_window_s
                ):
                    lift_onset_time = None
                    lift_streak = 0
                    lift_condition_times.clear()

        event: EventRecord | None = None
        if touch_condition:
            if mode == "onset" and touch_onset_time is None:
                touch_onset_time = frame.timestamp
            touch_streak += 1
            touch_condition_times.append(frame.timestamp)
            lift_streak = _decay_streak(lift_streak, lift_condition_times)
            if touch_streak >= config.confirm_samples:
                first_time = touch_condition_times[0]
                boundary_time: float | None = None
                fallback_reason = ""
                boundary_track = _find_raw_track(raw_tracks, primary_track_id)
                if boundary_track is None:
                    fallback_reason = "no_associated_track"
                elif (
                    config.raw_track_max_touch_candidate_age_s is not None
                    and first_time - boundary_track.first_seen_time
                    > config.raw_track_max_touch_candidate_age_s
                ):
                    fallback_reason = "touch_candidate_too_old"
                else:
                    boundary_time = boundary_track.first_seen_time
                event_time = _event_time_for_mode(
                    mode,
                    "touch",
                    frame.timestamp,
                    first_time,
                    touch_onset_time,
                    boundary_time,
                    config.confirm_window_s,
                )
                event = EventRecord(
                    kind="touch",
                    time=event_time,
                    ratio=stats.ratio,
                    centroid_cm=stats.centroid_cm,
                    frame_index=frame.index,
                    confirm_time=frame.timestamp,
                    first_confirm_frame_time=first_time,
                    onset_time=touch_onset_time,
                    associated_track_id=primary_track_id,
                    boundary_time=boundary_time,
                    fallback_reason=fallback_reason,
                )
                events.append(event)
                state = "ground"
                active_contact_track_id = primary_track_id
                touch_streak = 0
                touch_condition_times.clear()
                touch_onset_time = None
                lift_onset_time = None
        elif lift_condition:
            if mode == "onset" and lift_onset_time is None:
                lift_onset_time = frame.timestamp
            lift_streak += 1
            lift_condition_times.append(frame.timestamp)
            touch_streak = _decay_streak(touch_streak, touch_condition_times)
            if lift_streak >= config.confirm_samples:
                first_time = lift_condition_times[0]
                boundary_time = None
                fallback_reason = ""
                boundary_track = _find_raw_track(raw_tracks, active_contact_track_id)
                if boundary_track is None:
                    fallback_reason = "no_active_contact_track"
                else:
                    boundary_time = boundary_track.last_seen_time
                event_time = _event_time_for_mode(
                    mode,
                    "lift",
                    frame.timestamp,
                    first_time,
                    lift_onset_time,
                    boundary_time,
                    config.confirm_window_s,
                )
                event = EventRecord(
                    kind="lift",
                    time=event_time,
                    ratio=stats.ratio,
                    centroid_cm=stats.centroid_cm,
                    frame_index=frame.index,
                    confirm_time=frame.timestamp,
                    first_confirm_frame_time=first_time,
                    onset_time=lift_onset_time,
                    associated_track_id=active_contact_track_id,
                    boundary_time=boundary_time,
                    fallback_reason=fallback_reason,
                )
                events.append(event)
                state = "air"
                active_contact_track_id = None
                lift_streak = 0
                lift_condition_times.clear()
                lift_onset_time = None
                touch_onset_time = None
        else:
            touch_streak = _decay_streak(touch_streak, touch_condition_times)
            lift_streak = _decay_streak(lift_streak, lift_condition_times)

        traces.append(
            FrameTrace(
                variant=mode,
                confirm_window_s=config.confirm_window_s,
                frame_index=frame.index,
                timestamp=frame.timestamp,
                active_led_count=stats.active_led_count,
                raw_primary_cluster_length=stats.raw_primary_cluster_length,
                valid_primary_cluster_length=stats.valid_primary_cluster_length,
                ratio=stats.ratio,
                state=state_before,
                touch_condition=touch_condition,
                lift_condition=lift_condition,
                touch_streak=touch_streak,
                lift_streak=lift_streak,
                event_kind=event.kind if event else "",
                event_time_current=event.confirm_time if event else None,
                event_time_first_confirm_frame=(
                    event.first_confirm_frame_time if event else None
                ),
                onset_time=event.onset_time if event else None,
                confirm_time=event.confirm_time if event else None,
                associated_track_id=event.associated_track_id if event else None,
                boundary_time=event.boundary_time if event else None,
                fallback_reason=event.fallback_reason if event else "",
            )
        )

    return DetectorResult(events=events, traces=traces)


def compute_jump_rows(events: list[EventRecord], pairing_mode: PairingMode) -> list[JumpRow]:
    if pairing_mode == "raw_triplet":
        touches = [event for event in events if event.kind == "touch"]
        lifts = [event for event in events if event.kind == "lift"]
        rows: list[JumpRow] = []
        for index, lift in enumerate(lifts):
            if index >= len(touches) - 1:
                break
            touch = touches[index]
            next_touch = touches[index + 1]
            if not (touch.time < lift.time < next_touch.time):
                continue
            rows.append(
                JumpRow(
                    jump_index=len(rows) + 1,
                    touch_start_s=touch.time,
                    lift_s=lift.time,
                    landing_s=next_touch.time,
                    contact_time_s=lift.time - touch.time,
                    air_time_s=next_touch.time - lift.time,
                )
            )
        return rows

    if pairing_mode != "production_equivalent":
        raise ValueError(f"Unknown pairing mode: {pairing_mode}")

    rows: list[JumpRow] = []
    touch_count = 0
    lift_count = 0
    last_touch_time: float | None = None
    last_lift_time: float | None = None
    pending_air: list[tuple[float, float, float]] = []
    pending_contact: list[tuple[float, float, float]] = []

    for event in events:
        if event.kind == "touch":
            touch_count += 1
            if last_lift_time is not None:
                air_time = event.time - last_lift_time
                if air_time > 0:
                    pending_air.append((last_lift_time, event.time, air_time))
            last_touch_time = event.time
        elif event.kind == "lift":
            lift_count += 1
            if last_touch_time is not None and lift_count > 1:
                contact_time = event.time - last_touch_time
                if contact_time > 0:
                    pending_contact.append((last_touch_time, event.time, contact_time))
            last_lift_time = event.time

    max_len = max(len(pending_air), len(pending_contact))
    for index in range(max_len):
        air = pending_air[index] if index < len(pending_air) else None
        contact = pending_contact[index] if index < len(pending_contact) else None
        rows.append(
            JumpRow(
                jump_index=index + 1,
                touch_start_s=contact[0] if contact else None,
                lift_s=contact[1] if contact else (air[0] if air else None),
                landing_s=air[1] if air else None,
                contact_time_s=contact[2] if contact else None,
                air_time_s=air[2] if air else None,
            )
        )
    return rows


def read_manual_rows(path: str | Path) -> list[ManualRow]:
    rows: list[ManualRow] = []
    with open(path, "r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            jump_index = int(row["jump_index"])
            if row.get("manual_contact_time_s") and row.get("manual_air_time_s"):
                rows.append(
                    ManualRow(
                        jump_index=jump_index,
                        contact_time_s=float(row["manual_contact_time_s"]),
                        air_time_s=float(row["manual_air_time_s"]),
                    )
                )
                continue
            rows.append(
                ManualRow(
                    jump_index=jump_index,
                    contact_time_s=float(row["manual_lift_s"])
                    - float(row["manual_touch_start_s"]),
                    air_time_s=float(row["manual_next_touch_s"])
                    - float(row["manual_lift_s"]),
                )
            )
    return rows


def compare_rows(
    manual_rows: list[ManualRow],
    algorithm_rows: list[JumpRow],
    variant: str,
    pairing_mode: PairingMode,
    confirm_window_s: float,
) -> list[dict[str, object]]:
    comparisons: list[dict[str, object]] = []
    algorithm_by_index = {row.jump_index: row for row in algorithm_rows}
    for manual in manual_rows:
        algorithm = algorithm_by_index.get(manual.jump_index)
        contact = algorithm.contact_time_s if algorithm else None
        air = algorithm.air_time_s if algorithm else None
        comparisons.append(
            {
                "jump_index": manual.jump_index,
                "variant": variant,
                "pairing_mode": pairing_mode,
                "confirm_window_s": confirm_window_s,
                "manual_contact_time": manual.contact_time_s,
                "algorithm_contact_time": contact,
                "contact_error": contact - manual.contact_time_s
                if contact is not None
                else None,
                "manual_air_time": manual.air_time_s,
                "algorithm_air_time": air,
                "air_error": air - manual.air_time_s if air is not None else None,
            }
        )
    return comparisons


def _mean_absolute(values: Iterable[float | None]) -> float | None:
    concrete = [abs(value) for value in values if value is not None]
    if not concrete:
        return None
    return sum(concrete) / len(concrete)


def _fmt(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return f"{value:.9f}"
    return str(value)


def write_trace(path: Path, traces: list[FrameTrace]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = list(FrameTrace.__dataclass_fields__)
    with open(path, "w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        for trace in traces:
            writer.writerow({field: getattr(trace, field) for field in fields})


def write_comparisons(path: Path, rows: list[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = [
        "jump_index",
        "variant",
        "pairing_mode",
        "confirm_window_s",
        "manual_contact_time",
        "algorithm_contact_time",
        "contact_error",
        "manual_air_time",
        "algorithm_air_time",
        "air_error",
    ]
    with open(path, "w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        for row in rows:
            writer.writerow({field: _fmt(row.get(field)) for field in fields})


def write_summary(path: Path, comparisons: list[dict[str, object]]) -> None:
    grouped: dict[tuple[str, str, float], list[dict[str, object]]] = {}
    for row in comparisons:
        key = (
            str(row["variant"]),
            str(row["pairing_mode"]),
            float(row["confirm_window_s"]),
        )
        grouped.setdefault(key, []).append(row)

    lines = [
        "# Session 5 Jump Timing Diagnostics",
        "",
        "This report compares manual timing against offline detector variants.",
        "",
        "## Aggregate Error",
        "",
        "| variant | pairing_mode | confirm_window_s | contact_n | air_n | contact_MAE_s | air_MAE_s | contact_bias | air_bias |",
        "|---|---|---:|---:|---:|---:|---:|---|---|",
    ]
    aggregates: dict[tuple[str, str, float], dict[str, object]] = {}
    for key in sorted(grouped):
        variant, pairing_mode, window = key
        rows = grouped[key]
        contact_errors = [row["contact_error"] for row in rows]
        air_errors = [row["air_error"] for row in rows]
        contact_n = sum(1 for value in contact_errors if value is not None)
        air_n = sum(1 for value in air_errors if value is not None)
        contact_mae = _mean_absolute(contact_errors)
        air_mae = _mean_absolute(air_errors)
        contact_sum = sum(value for value in contact_errors if value is not None)
        air_sum = sum(value for value in air_errors if value is not None)
        contact_bias = "longer" if contact_sum > 0 else "shorter" if contact_sum < 0 else "even"
        air_bias = "longer" if air_sum > 0 else "shorter" if air_sum < 0 else "even"
        aggregates[key] = {
            "contact_n": contact_n,
            "air_n": air_n,
            "contact_mae": contact_mae,
            "air_mae": air_mae,
            "contact_bias": contact_bias,
            "air_bias": air_bias,
        }
        lines.append(
            "| "
            + " | ".join(
                [
                    variant,
                    pairing_mode,
                    f"{window:.2f}",
                    str(contact_n),
                    str(air_n),
                    _fmt(contact_mae),
                    _fmt(air_mae),
                    contact_bias,
                    air_bias,
                ]
            )
            + " |"
        )

    current_prod = aggregates.get(("current", "production_equivalent", 0.10))
    compensated_prod = aggregates.get(("compensated", "production_equivalent", 0.10))
    onset_prod = {
        key: value
        for key, value in aggregates.items()
        if key[0] == "onset" and key[1] == "production_equivalent"
    }
    associated_prod = {
        key: value
        for key, value in aggregates.items()
        if key[0] in {"touch_associated", "lift_associated", "associated"}
        and key[1] == "production_equivalent"
    }
    best_onset_key = None
    best_onset = None
    if onset_prod:
        best_onset_key = min(
            onset_prod,
            key=lambda key: (
                (onset_prod[key]["contact_mae"] or math.inf)
                + (onset_prod[key]["air_mae"] or math.inf)
            ),
        )
        best_onset = onset_prod[best_onset_key]
    best_associated_key = None
    best_associated = None
    if associated_prod:
        best_associated_key = min(
            associated_prod,
            key=lambda key: (
                (associated_prod[key]["contact_mae"] or math.inf)
                + (associated_prod[key]["air_mae"] or math.inf)
            ),
        )
        best_associated = associated_prod[best_associated_key]

    lines.extend(["", "## Findings", ""])
    if current_prod:
        lines.append(
            "- Current production-equivalent timing makes contact_time "
            f"{current_prod['contact_bias']} (MAE {_fmt(current_prod['contact_mae'])}s) "
            f"and air_time {current_prod['air_bias']} "
            f"(MAE {_fmt(current_prod['air_mae'])}s)."
        )
    if current_prod and compensated_prod:
        current_total = (current_prod["contact_mae"] or 0.0) + (
            current_prod["air_mae"] or 0.0
        )
        compensated_total = (compensated_prod["contact_mae"] or 0.0) + (
            compensated_prod["air_mae"] or 0.0
        )
        if compensated_total < current_total * 0.95:
            verdict = "materially improves"
        elif compensated_total > current_total * 1.05:
            verdict = "is materially worse than"
        else:
            verdict = "is effectively the same as"
        lines.append(f"- Compensated timing {verdict} current timing.")
    if best_onset:
        lines.append(
            "- Best onset+confirm production-equivalent window is "
            f"{best_onset_key[2]:.2f}s: contact MAE {_fmt(best_onset['contact_mae'])}s, "
            f"air MAE {_fmt(best_onset['air_mae'])}s."
        )
    if best_associated:
        lines.append(
            "- Best associated-track production-equivalent variant is "
            f"{best_associated_key[0]}: contact MAE "
            f"{_fmt(best_associated['contact_mae'])}s, air MAE "
            f"{_fmt(best_associated['air_mae'])}s."
        )
    lines.append(
        "- raw_triplet includes the initial stance contact period; production_equivalent "
        "matches the current GaitEngine reporting rules better."
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def run_diagnostics(
    frames_path: Path,
    manual_path: Path,
    out_dir: Path,
    confirm_windows: list[float],
) -> None:
    frames = list(iter_excel_led_frames(frames_path))
    manual_rows = read_manual_rows(manual_path)
    comparisons: list[dict[str, object]] = []
    traces: list[FrameTrace] = []

    for mode in ("current", "compensated"):
        config = TimingConfig()
        result = run_detector(frames, mode, config)
        traces.extend(result.traces)
        for pairing_mode in ("raw_triplet", "production_equivalent"):
            comparisons.extend(
                compare_rows(
                    manual_rows,
                    compute_jump_rows(result.events, pairing_mode),
                    mode,
                    pairing_mode,
                    config.confirm_window_s,
                )
            )

    for mode in ("touch_associated", "lift_associated", "associated"):
        config = TimingConfig()
        result = run_detector(frames, mode, config)
        traces.extend(result.traces)
        for pairing_mode in ("raw_triplet", "production_equivalent"):
            comparisons.extend(
                compare_rows(
                    manual_rows,
                    compute_jump_rows(result.events, pairing_mode),
                    mode,
                    pairing_mode,
                    config.confirm_window_s,
                )
            )

    for window in confirm_windows:
        config = TimingConfig(confirm_window_s=window)
        result = run_detector(frames, "onset", config)
        traces.extend(result.traces)
        for pairing_mode in ("raw_triplet", "production_equivalent"):
            comparisons.extend(
                compare_rows(
                    manual_rows,
                    compute_jump_rows(result.events, pairing_mode),
                    "onset",
                    pairing_mode,
                    window,
                )
            )

    write_trace(out_dir / "frame_trace.csv", traces)
    write_comparisons(out_dir / "event_compare.csv", comparisons)
    write_summary(out_dir / "summary.md", comparisons)


def _parse_windows(value: str) -> list[float]:
    return [float(part.strip()) for part in value.split(",") if part.strip()]


def main() -> int:
    parser = argparse.ArgumentParser(description="Offline jump timing diagnostics")
    parser.add_argument("--frames", required=True, type=Path)
    parser.add_argument("--manual", required=True, type=Path)
    parser.add_argument("--out-dir", required=True, type=Path)
    parser.add_argument(
        "--sweep-confirm-windows",
        default="0.05,0.10,0.15,0.25",
        help="Comma-separated onset confirmation windows in seconds",
    )
    args = parser.parse_args()
    run_diagnostics(
        args.frames,
        args.manual,
        args.out_dir,
        _parse_windows(args.sweep_confirm_windows),
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
